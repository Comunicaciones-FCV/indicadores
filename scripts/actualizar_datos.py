#!/usr/bin/env python3
"""Extrae la pestaña "Comunicaciones" del Excel de indicadores FCV y genera
datos/comunicaciones.js, que el tablero (index.html) lee al abrirse.

Uso:
    python3 scripts/actualizar_datos.py "INDICADORES ÁREAS FCV 2026.xlsx"

Requiere: pip install openpyxl
"""
import datetime
import json
import pathlib
import sys

import openpyxl

RAIZ = pathlib.Path(__file__).resolve().parent.parent
SALIDA = RAIZ / "datos" / "comunicaciones.js"


def main():
    if len(sys.argv) < 2:
        sys.exit("Indica la ruta del Excel: python3 scripts/actualizar_datos.py archivo.xlsx")
    origen = pathlib.Path(sys.argv[1])
    libro = openpyxl.load_workbook(origen, data_only=True)
    hojas = [h for h in libro.worksheets if "comunic" in h.title.lower()]
    if not hojas:
        sys.exit(f"No se encontró una pestaña 'Comunicaciones' en {origen.name}: {libro.sheetnames}")
    hoja = hojas[0]

    celdas = []
    for fila in hoja.iter_rows(values_only=True):
        celdas.append([v if v is None or isinstance(v, (int, float)) else str(v) for v in fila])
    # Quita filas vacías del final
    while celdas and all(v is None for v in celdas[-1]):
        celdas.pop()

    datos = {
        "fuente": origen.name,
        "hoja": hoja.title,
        "actualizado": datetime.date.today().isoformat(),
        "celdas": celdas,
    }
    filas = ",\n  ".join(json.dumps(f, ensure_ascii=False) for f in celdas)
    cabecera = {k: v for k, v in datos.items() if k != "celdas"}
    SALIDA.write_text(
        "// Generado por scripts/actualizar_datos.py — no editar a mano.\n"
        "window.DATOS_FCV = " + json.dumps(cabecera, ensure_ascii=False)[:-1]
        + ', "celdas": [\n  ' + filas + "\n]};\n",
        encoding="utf-8",
    )
    print(f"OK: {SALIDA.relative_to(RAIZ)} ({len(celdas)} filas de '{hoja.title}')")


if __name__ == "__main__":
    main()
